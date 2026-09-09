from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import msoffcrypto
import pytest
from openpyxl import Workbook, load_workbook

from enam_assessment.errors import (
    MissingInputError,
    WorkbookPasswordError,
    WorkbookReadError,
    WorkbookSchemaError,
)
from enam_assessment.ingestion import (
    LotClassification,
    ValidationSeverity,
    ingest_workbook,
    render_ingestion_report,
)

EXPECTED_SHEETS = ("Amber", "DBL", "Zee", "Welspun")
HEADERS = ("Date", "Qty", "Rate", "Amount", "Date", "Qty", "Rate", "Amount")


def write_synthetic_workbook(
    path: Path,
    rows_by_sheet: dict[str, list[tuple[Any, ...]]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_name in EXPECTED_SHEETS:
        sheet = workbook.create_sheet(sheet_name)
        sheet.cell(1, 1, f"Synthetic {sheet_name}")
        sheet.cell(2, 1, "Purchase")
        sheet.cell(2, 5, "Sale")
        for column, header in enumerate(HEADERS, start=1):
            sheet.cell(3, column, header)
        for row in rows_by_sheet.get(sheet_name, []):
            sheet.append(row)

    workbook.save(path)
    workbook.close()


def encrypt_workbook(source_path: Path, encrypted_path: Path, password: str) -> None:
    with source_path.open("rb") as source, encrypted_path.open("wb") as destination:
        msoffcrypto.OfficeFile(source).encrypt(password, destination)


def test_ingests_realized_and_open_lots_with_source_lineage(tmp_path: Path) -> None:
    workbook_path = tmp_path / "trades.xlsx"
    write_synthetic_workbook(
        workbook_path,
        {
            "Amber": [
                (
                    datetime(2024, 1, 2),
                    10,
                    100.25,
                    1002.50,
                    datetime(2024, 2, 2),
                    10,
                    110.50,
                    1105.00,
                ),
                (datetime(2024, 3, 5), 4, 95, 380, None, None, None, None),
            ]
        },
    )

    result = ingest_workbook(workbook_path, password="synthetic-password")
    amber_rows = result.records_for("Amber")

    assert amber_rows[0].classification is LotClassification.IGNORED_STRUCTURAL
    assert amber_rows[0].source_sheet == "Amber"
    assert amber_rows[0].source_row == 1

    realized = amber_rows[3]
    assert realized.classification is LotClassification.VALID_REALIZED
    assert realized.company == "Amber Enterprises"
    assert realized.purchase_quantity == Decimal("10")
    assert realized.sale_amount == Decimal("1105")
    assert realized.source_sheet == "Amber"
    assert realized.source_row == 4
    assert realized.validation_results == ()

    open_lot = amber_rows[4]
    assert open_lot.classification is LotClassification.PROVISIONALLY_OPEN
    assert open_lot.purchase_amount == Decimal("380")
    assert open_lot.sale_date is None
    assert open_lot.source_row == 5


def test_classifies_invalid_structural_and_post_period_rows(tmp_path: Path) -> None:
    workbook_path = tmp_path / "validation-cases.xlsx"
    write_synthetic_workbook(
        workbook_path,
        {
            "Amber": [
                ("not-a-date", 10, 100, 1000, None, None, None, None),
                (datetime(2024, 1, 2), 10, 100, 1000.03, None, None, None, None),
                (
                    datetime(2024, 2, 2),
                    5,
                    100,
                    500,
                    datetime(2024, 1, 2),
                    5,
                    110,
                    550,
                ),
                (datetime(2024, 1, 2), 5, 100, 500, datetime(2024, 2, 2), 5, None, 550),
                (None, None, None, None, None, None, None, None),
                (None, 15, None, 1500, None, 15, None, 1700),
                (
                    datetime(2025, 4, 2),
                    2,
                    100,
                    200,
                    datetime(2025, 4, 3),
                    2,
                    110,
                    220,
                ),
            ]
        },
    )

    result = ingest_workbook(workbook_path, password="synthetic-password")
    rows = result.records_for("Amber")

    expected = [
        (4, LotClassification.QUARANTINED_INVALID, {"invalid_purchase_date"}),
        (5, LotClassification.QUARANTINED_INVALID, {"purchase_amount_mismatch"}),
        (6, LotClassification.QUARANTINED_INVALID, {"purchase_after_sale"}),
        (7, LotClassification.QUARANTINED_INVALID, {"partial_sale"}),
        (8, LotClassification.IGNORED_STRUCTURAL, {"blank"}),
        (9, LotClassification.IGNORED_STRUCTURAL, {"undated_aggregate"}),
        (10, LotClassification.VALID_REALIZED, {"after_stated_period"}),
    ]
    for row_number, classification, expected_codes in expected:
        record = rows[row_number - 1]
        assert record.source_row == row_number
        assert record.classification is classification
        assert {result.code for result in record.validation_results} == expected_codes

    assert rows[9].validation_results[0].severity is ValidationSeverity.WARNING


def test_requires_readable_file_and_nonempty_password(tmp_path: Path) -> None:
    with pytest.raises(MissingInputError, match="Workbook file does not exist"):
        ingest_workbook(tmp_path / "missing.xlsx", password="provided")

    workbook_path = tmp_path / "plain.xlsx"
    write_synthetic_workbook(workbook_path, {})
    with pytest.raises(WorkbookPasswordError, match="password is required"):
        ingest_workbook(workbook_path, password="")

    unreadable_path = tmp_path / "unreadable.xlsx"
    unreadable_path.write_text("not an Excel workbook", encoding="utf-8")
    with pytest.raises(WorkbookReadError, match="could not be read"):
        ingest_workbook(unreadable_path, password="provided")


def test_decrypts_in_memory_and_rejects_wrong_password(tmp_path: Path) -> None:
    plain_path = tmp_path / "plain.xlsx"
    encrypted_path = tmp_path / "encrypted.xlsx"
    write_synthetic_workbook(
        plain_path,
        {"Amber": [(datetime(2024, 1, 2), 1, 100, 100, None, None, None, None)]},
    )
    encrypt_workbook(plain_path, encrypted_path, "correct-password")

    result = ingest_workbook(encrypted_path, password="correct-password")
    assert result.records_for("Amber")[3].classification is LotClassification.PROVISIONALLY_OPEN
    with pytest.raises(WorkbookPasswordError, match="incorrect"):
        ingest_workbook(encrypted_path, password="wrong-password")


def test_rejects_missing_sheet_and_changed_columns(tmp_path: Path) -> None:
    missing_sheet_path = tmp_path / "missing-sheet.xlsx"
    write_synthetic_workbook(missing_sheet_path, {})
    workbook = load_workbook(missing_sheet_path)
    del workbook["Zee"]
    workbook.save(missing_sheet_path)
    workbook.close()

    with pytest.raises(WorkbookSchemaError, match="missing required sheets.*Zee"):
        ingest_workbook(missing_sheet_path, password="provided")

    changed_columns_path = tmp_path / "changed-columns.xlsx"
    write_synthetic_workbook(changed_columns_path, {})
    workbook = load_workbook(changed_columns_path)
    workbook["Amber"].cell(3, 3, "Price")
    workbook.save(changed_columns_path)
    workbook.close()

    with pytest.raises(WorkbookSchemaError, match="Amber.*expected columns"):
        ingest_workbook(changed_columns_path, password="provided")


def test_enforces_numeric_fields_amount_tolerance_and_preserves_repeats(tmp_path: Path) -> None:
    workbook_path = tmp_path / "numeric-cases.xlsx"
    repeated_row = (datetime(2024, 1, 2), 10, 100, 1000.02, None, None, None, None)
    write_synthetic_workbook(
        workbook_path,
        {
            "Amber": [
                repeated_row,
                repeated_row,
                (
                    datetime(2024, 1, 2),
                    5,
                    100,
                    500,
                    datetime(2024, 2, 2),
                    5,
                    110,
                    550.03,
                ),
                (datetime(2024, 1, 2), 0, 100, 0, None, None, None, None),
                (datetime(2024, 1, 2), 1, "100", 100, None, None, None, None),
                (datetime(2024, 1, 2), 1, None, 100, None, None, None, None),
            ]
        },
    )

    rows = ingest_workbook(workbook_path, password="provided").records_for("Amber")

    assert rows[3].classification is LotClassification.PROVISIONALLY_OPEN
    assert rows[4].classification is LotClassification.PROVISIONALLY_OPEN
    assert rows[3].purchase_amount == rows[4].purchase_amount == Decimal("1000.02")
    assert {result.code for result in rows[5].validation_results} == {"sale_amount_mismatch"}
    assert "nonpositive_purchase_quantity" in {result.code for result in rows[6].validation_results}
    assert "invalid_purchase_rate" in {result.code for result in rows[7].validation_results}
    assert "missing_purchase_rate" in {result.code for result in rows[8].validation_results}


def test_quarantines_sale_without_purchase_and_keeps_ambiguity_unresolved(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "sale-only.xlsx"
    write_synthetic_workbook(
        workbook_path,
        {
            "Amber": [
                (
                    None,
                    None,
                    None,
                    None,
                    datetime(2024, 2, 2),
                    5,
                    110,
                    550,
                )
            ]
        },
    )

    record = ingest_workbook(workbook_path, password="provided").records_for("Amber")[3]

    assert record.classification is LotClassification.QUARANTINED_INVALID
    assert "sale_without_purchase" in {validation.code for validation in record.validation_results}


def test_ignores_far_right_formula_when_finding_relevant_rows(tmp_path: Path) -> None:
    workbook_path = tmp_path / "far-right.xlsx"
    write_synthetic_workbook(
        workbook_path,
        {"Welspun": [(datetime(2024, 1, 2), 2, 50, 100, None, None, None, None)]},
    )
    workbook = load_workbook(workbook_path)
    workbook["Welspun"].cell(100, 16377, "=SUM(XEW90:XFD99)")
    workbook.save(workbook_path)
    workbook.close()

    rows = ingest_workbook(workbook_path, password="provided").records_for("Welspun")

    assert len(rows) == 4
    assert rows[-1].source_row == 4
    assert rows[-1].classification is LotClassification.PROVISIONALLY_OPEN


def test_reconciles_accepted_lots_and_renders_reason_counts(tmp_path: Path) -> None:
    workbook_path = tmp_path / "reconciliation.xlsx"
    write_synthetic_workbook(
        workbook_path,
        {
            "Amber": [
                (
                    datetime(2024, 1, 2),
                    10,
                    100,
                    1000,
                    datetime(2024, 2, 2),
                    10,
                    110,
                    1100,
                ),
                (datetime(2024, 3, 5), 4, 95, 380, None, None, None, None),
                (datetime(2024, 4, 5), 1, 100, 100.03, None, None, None, None),
                (None, None, None, None, None, None, None, None),
                (None, 15, None, 1500, None, 15, None, 1700),
            ]
        },
    )

    result = ingest_workbook(workbook_path, password="provided")
    reconciliation = result.reconciliation_for("Amber")

    assert reconciliation.source_rows_considered == 5
    assert reconciliation.valid_realized_lots == 1
    assert reconciliation.provisionally_open_lots == 1
    assert reconciliation.quarantined_rows == 1
    assert reconciliation.ignored_rows == 2
    assert reconciliation.purchase_quantity == Decimal("14")
    assert reconciliation.sale_quantity == Decimal("10")
    assert reconciliation.open_quantity == Decimal("4")
    assert reconciliation.realized_quantity_difference == Decimal("0")
    assert reconciliation.purchase_amount == Decimal("1380")
    assert reconciliation.sale_amount == Decimal("1100")
    assert reconciliation.purchase_amount_difference == Decimal("0")
    assert reconciliation.sale_amount_difference == Decimal("0")
    assert reconciliation.earliest_transaction_date.isoformat() == "2024-01-02"
    assert reconciliation.latest_transaction_date.isoformat() == "2024-03-05"
    assert reconciliation.reason_counts == (
        ("blank", 1),
        ("purchase_amount_mismatch", 1),
        ("undated_aggregate", 1),
    )
    assert reconciliation.reason_rows == (
        ("blank", (7,)),
        ("purchase_amount_mismatch", (6,)),
        ("undated_aggregate", (8,)),
    )

    report = render_ingestion_report(result, workbook_name="synthetic.xlsx")
    assert "# Historical Trade Workbook Ingestion Report" in report
    assert "| Amber | 5 | 1 | 1 | 1 | 2 |" in report
    assert "purchase_amount_mismatch: 1" in report
    assert "`purchase_amount_mismatch`: rows 6" in report
    assert "No portfolio NAV or execution-ledger completeness is claimed" in report


def test_report_normalizes_insignificant_negative_amount_difference(tmp_path: Path) -> None:
    workbook_path = tmp_path / "negative-zero.xlsx"
    write_synthetic_workbook(
        workbook_path,
        {
            "Amber": [
                (
                    datetime(2024, 1, 2),
                    3,
                    0.33333333333333337,
                    1,
                    datetime(2024, 2, 2),
                    3,
                    0.33333333333333337,
                    1,
                )
            ]
        },
    )

    result = ingest_workbook(workbook_path, password="provided")
    report = render_ingestion_report(result, workbook_name="synthetic.xlsx")

    assert "-0.00" not in report
