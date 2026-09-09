"""Read-only ingestion of the assessment's FIFO-matched trade-lot workbook."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from enum import StrEnum
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import msoffcrypto
from msoffcrypto.exceptions import DecryptionError, FileFormatError, InvalidKeyError, ParseError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .errors import (
    MissingInputError,
    WorkbookPasswordError,
    WorkbookReadError,
    WorkbookSchemaError,
)

AMOUNT_TOLERANCE = Decimal("0.02")
STATED_PERIOD_END = date(2025, 3, 31)


class LotClassification(StrEnum):
    """Mutually exclusive row outcomes produced by ingestion."""

    VALID_REALIZED = "valid_realized"
    PROVISIONALLY_OPEN = "provisionally_open"
    QUARANTINED_INVALID = "quarantined_invalid"
    IGNORED_STRUCTURAL = "ignored_structural"


class ValidationSeverity(StrEnum):
    """Severity of a row-level validation observation."""

    INFO = "info"
    WARNING = "warning"
    ERROR = "error"


@dataclass(frozen=True, slots=True)
class ValidationResult:
    """One explicit validation observation for a source row."""

    code: str
    message: str
    severity: ValidationSeverity


@dataclass(frozen=True, slots=True)
class TradeLotRecord:
    """Normalized values and lineage for one source row."""

    company: str
    purchase_date: date | None
    purchase_quantity: Decimal | None
    purchase_rate: Decimal | None
    purchase_amount: Decimal | None
    sale_date: date | None
    sale_quantity: Decimal | None
    sale_rate: Decimal | None
    sale_amount: Decimal | None
    classification: LotClassification
    source_sheet: str
    source_row: int
    validation_results: tuple[ValidationResult, ...]


@dataclass(frozen=True, slots=True)
class IngestionResult:
    """All classified rows from the four expected sheets."""

    records: tuple[TradeLotRecord, ...]

    def records_for(self, sheet_name: str) -> tuple[TradeLotRecord, ...]:
        return tuple(record for record in self.records if record.source_sheet == sheet_name)

    def reconciliation_for(self, sheet_name: str) -> SheetReconciliation:
        if sheet_name not in _COMPANIES:
            raise KeyError(f"Unknown trade sheet: {sheet_name}")
        return _reconcile(self.records_for(sheet_name), sheet_name, _COMPANIES[sheet_name])


@dataclass(frozen=True, slots=True)
class SheetReconciliation:
    """Accepted-lot totals and complete classification counts for one sheet."""

    source_sheet: str
    company: str
    source_rows_considered: int
    valid_realized_lots: int
    provisionally_open_lots: int
    quarantined_rows: int
    ignored_rows: int
    purchase_quantity: Decimal
    sale_quantity: Decimal
    open_quantity: Decimal
    realized_quantity_difference: Decimal
    purchase_amount: Decimal
    sale_amount: Decimal
    purchase_amount_difference: Decimal
    sale_amount_difference: Decimal
    earliest_transaction_date: date | None
    latest_transaction_date: date | None
    after_stated_period_rows: int
    reason_counts: tuple[tuple[str, int], ...]
    reason_rows: tuple[tuple[str, tuple[int, ...]], ...]


_COMPANIES = {
    "Amber": "Amber Enterprises",
    "DBL": "Dilip Buildcon",
    "Zee": "Zee Entertainment Enterprises",
    "Welspun": "Welspun Living",
}
_EXPECTED_HEADERS = ("Date", "Qty", "Rate", "Amount", "Date", "Qty", "Rate", "Amount")


def ingest_workbook(path: Path, password: str) -> IngestionResult:
    """Load an encrypted or plain workbook without writing decrypted bytes to disk."""
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise MissingInputError(f"Workbook file does not exist: {workbook_path}")
    if not password:
        raise WorkbookPasswordError("A non-empty workbook password is required.")

    stream = _workbook_stream(workbook_path, password)
    try:
        workbook = load_workbook(stream, read_only=False, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as error:
        raise WorkbookReadError(
            f"Workbook could not be read as a valid .xlsx file: {workbook_path}"
        ) from error
    try:
        _validate_schema(workbook.sheetnames, workbook)
        records: list[TradeLotRecord] = []
        for sheet_name, company in _COMPANIES.items():
            sheet = workbook[sheet_name]
            last_row = _last_relevant_row(sheet)
            for row_number in range(1, last_row + 1):
                values = tuple(sheet.cell(row_number, column).value for column in range(1, 9))
                records.append(_normalize_row(company, sheet_name, row_number, values))
        return IngestionResult(records=tuple(records))
    finally:
        workbook.close()


def _workbook_stream(path: Path, password: str) -> BytesIO:
    try:
        with path.open("rb") as source:
            office_file = msoffcrypto.OfficeFile(source)
            if office_file.is_encrypted():
                try:
                    office_file.load_key(password=password, verify_password=True)
                    output = BytesIO()
                    office_file.decrypt(output, verify_integrity=True)
                except (DecryptionError, InvalidKeyError) as error:
                    raise WorkbookPasswordError(
                        "Workbook password is incorrect or cannot decrypt this workbook."
                    ) from error
                output.seek(0)
                return output
            source.seek(0)
            return BytesIO(source.read())
    except WorkbookPasswordError:
        raise
    except (FileFormatError, ParseError, OSError) as error:
        raise WorkbookReadError(f"Workbook could not be read: {path}") from error


def _validate_schema(sheet_names: list[str], workbook: Any) -> None:
    missing_sheets = [sheet for sheet in _COMPANIES if sheet not in sheet_names]
    if missing_sheets:
        raise WorkbookSchemaError(
            "Workbook is missing required sheets: " + ", ".join(missing_sheets)
        )

    for sheet_name in _COMPANIES:
        sheet = workbook[sheet_name]
        headers = tuple(sheet.cell(3, column).value for column in range(1, 9))
        if headers != _EXPECTED_HEADERS:
            raise WorkbookSchemaError(
                f"Sheet {sheet_name} does not have the expected columns in A3:H3; "
                f"received {headers!r}."
            )


def _last_relevant_row(sheet: Any) -> int:
    for row_number in range(sheet.max_row, 0, -1):
        if any(sheet.cell(row_number, column).value is not None for column in range(1, 9)):
            return row_number
    return 0


def _normalize_row(
    company: str,
    sheet_name: str,
    row_number: int,
    values: tuple[Any, ...],
) -> TradeLotRecord:
    if row_number <= 3 or all(value is None for value in values):
        code = "header" if row_number <= 3 else "blank"
        return TradeLotRecord(
            company=company,
            purchase_date=None,
            purchase_quantity=None,
            purchase_rate=None,
            purchase_amount=None,
            sale_date=None,
            sale_quantity=None,
            sale_rate=None,
            sale_amount=None,
            classification=LotClassification.IGNORED_STRUCTURAL,
            source_sheet=sheet_name,
            source_row=row_number,
            validation_results=(
                ValidationResult(
                    code=code,
                    message="Workbook header row." if code == "header" else "Blank row.",
                    severity=ValidationSeverity.INFO,
                ),
            ),
        )

    if _is_undated_aggregate(values):
        return TradeLotRecord(
            company=company,
            purchase_date=None,
            purchase_quantity=_to_decimal(values[1]),
            purchase_rate=_to_decimal(values[2]),
            purchase_amount=_to_decimal(values[3]),
            sale_date=None,
            sale_quantity=_to_decimal(values[5]),
            sale_rate=_to_decimal(values[6]),
            sale_amount=_to_decimal(values[7]),
            classification=LotClassification.IGNORED_STRUCTURAL,
            source_sheet=sheet_name,
            source_row=row_number,
            validation_results=(
                ValidationResult(
                    code="undated_aggregate",
                    message="Undated numeric aggregate row; not a trade lot.",
                    severity=ValidationSeverity.INFO,
                ),
            ),
        )

    raw_purchase_date, raw_purchase_quantity, raw_purchase_rate, raw_purchase_amount = values[:4]
    raw_sale_date, raw_sale_quantity, raw_sale_rate, raw_sale_amount = values[4:]
    validation_results: list[ValidationResult] = []
    if all(value is None for value in values[:4]) and any(
        value is not None for value in values[4:]
    ):
        validation_results.append(
            ValidationResult(
                code="sale_without_purchase",
                message=(
                    "Sale details have no purchase-side lot. The workbook's unmatched-sell "
                    "meaning remains unresolved, so this row is quarantined."
                ),
                severity=ValidationSeverity.ERROR,
            )
        )

    purchase_date = _parse_date(
        raw_purchase_date, "purchase_date", validation_results, required=True
    )
    purchase_quantity = _parse_decimal(
        raw_purchase_quantity, "purchase_quantity", validation_results, required=True
    )
    purchase_rate = _parse_decimal(
        raw_purchase_rate, "purchase_rate", validation_results, required=True
    )
    purchase_amount = _parse_decimal(
        raw_purchase_amount, "purchase_amount", validation_results, required=True
    )

    has_sale = any(value is not None for value in values[4:])
    if has_sale and any(value is None for value in values[4:]):
        validation_results.append(
            ValidationResult(
                code="partial_sale",
                message="Sale details must provide date, quantity, rate, and amount together.",
                severity=ValidationSeverity.ERROR,
            )
        )
    sale_date = _parse_date(raw_sale_date, "sale_date", validation_results, required=False)
    sale_quantity = _parse_decimal(
        raw_sale_quantity, "sale_quantity", validation_results, required=False
    )
    sale_rate = _parse_decimal(raw_sale_rate, "sale_rate", validation_results, required=False)
    sale_amount = _parse_decimal(raw_sale_amount, "sale_amount", validation_results, required=False)

    _validate_amount(
        "purchase", purchase_quantity, purchase_rate, purchase_amount, validation_results
    )
    _validate_amount("sale", sale_quantity, sale_rate, sale_amount, validation_results)

    if purchase_date is not None and sale_date is not None and purchase_date > sale_date:
        validation_results.append(
            ValidationResult(
                code="purchase_after_sale",
                message="Purchase date occurs after sale date.",
                severity=ValidationSeverity.ERROR,
            )
        )

    dated_values = (value for value in (purchase_date, sale_date) if value is not None)
    if any(value > STATED_PERIOD_END for value in dated_values):
        validation_results.append(
            ValidationResult(
                code="after_stated_period",
                message="Transaction date is after the workbook's stated 2025-03-31 end date.",
                severity=ValidationSeverity.WARNING,
            )
        )

    has_errors = any(result.severity is ValidationSeverity.ERROR for result in validation_results)
    return TradeLotRecord(
        company=company,
        purchase_date=purchase_date,
        purchase_quantity=purchase_quantity,
        purchase_rate=purchase_rate,
        purchase_amount=purchase_amount,
        sale_date=sale_date,
        sale_quantity=sale_quantity,
        sale_rate=sale_rate,
        sale_amount=sale_amount,
        classification=(
            LotClassification.QUARANTINED_INVALID
            if has_errors
            else (
                LotClassification.VALID_REALIZED
                if has_sale
                else LotClassification.PROVISIONALLY_OPEN
            )
        ),
        source_sheet=sheet_name,
        source_row=row_number,
        validation_results=tuple(validation_results),
    )


def _is_undated_aggregate(values: tuple[Any, ...]) -> bool:
    if values[0] is not None or values[4] is not None:
        return False
    populated = [value for value in values[1:4] + values[5:8] if value is not None]
    return bool(populated) and all(_is_number(value) for value in populated)


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def _parse_date(
    value: Any,
    field: str,
    results: list[ValidationResult],
    *,
    required: bool,
) -> date | None:
    if value is None:
        if required:
            results.append(
                ValidationResult(
                    code=f"missing_{field}",
                    message=f"{field} is required.",
                    severity=ValidationSeverity.ERROR,
                )
            )
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    results.append(
        ValidationResult(
            code=f"invalid_{field}",
            message=f"{field} must be an Excel date; received {value!r}.",
            severity=ValidationSeverity.ERROR,
        )
    )
    return None


def _parse_decimal(
    value: Any,
    field: str,
    results: list[ValidationResult],
    *,
    required: bool,
) -> Decimal | None:
    if value is None:
        if required:
            results.append(
                ValidationResult(
                    code=f"missing_{field}",
                    message=f"{field} is required.",
                    severity=ValidationSeverity.ERROR,
                )
            )
        return None
    if not _is_number(value):
        results.append(
            ValidationResult(
                code=f"invalid_{field}",
                message=f"{field} must be numeric; received {value!r}.",
                severity=ValidationSeverity.ERROR,
            )
        )
        return None
    normalized = Decimal(str(value))
    if normalized <= 0:
        results.append(
            ValidationResult(
                code=f"nonpositive_{field}",
                message=f"{field} must be greater than zero; received {value!r}.",
                severity=ValidationSeverity.ERROR,
            )
        )
    return normalized


def _to_decimal(value: Any) -> Decimal | None:
    return Decimal(str(value)) if value is not None else None


def _validate_amount(
    side: str,
    quantity: Decimal | None,
    rate: Decimal | None,
    amount: Decimal | None,
    results: list[ValidationResult],
) -> None:
    if quantity is None or rate is None or amount is None:
        return
    difference = amount - (quantity * rate)
    if abs(difference) > AMOUNT_TOLERANCE:
        results.append(
            ValidationResult(
                code=f"{side}_amount_mismatch",
                message=(
                    f"{side}_amount differs from quantity multiplied by rate by "
                    f"{difference}; allowed absolute tolerance is {AMOUNT_TOLERANCE}."
                ),
                severity=ValidationSeverity.ERROR,
            )
        )


def _reconcile(
    records: tuple[TradeLotRecord, ...], sheet_name: str, company: str
) -> SheetReconciliation:
    data_records = tuple(record for record in records if record.source_row > 3)
    realized = tuple(
        record
        for record in data_records
        if record.classification is LotClassification.VALID_REALIZED
    )
    open_lots = tuple(
        record
        for record in data_records
        if record.classification is LotClassification.PROVISIONALLY_OPEN
    )
    quarantined = tuple(
        record
        for record in data_records
        if record.classification is LotClassification.QUARANTINED_INVALID
    )
    ignored = tuple(
        record
        for record in data_records
        if record.classification is LotClassification.IGNORED_STRUCTURAL
    )
    accepted = realized + open_lots

    purchase_quantity = _sum_field(accepted, "purchase_quantity")
    sale_quantity = _sum_field(realized, "sale_quantity")
    open_quantity = _sum_field(open_lots, "purchase_quantity")
    realized_purchase_quantity = _sum_field(realized, "purchase_quantity")
    purchase_amount = _sum_field(accepted, "purchase_amount")
    sale_amount = _sum_field(realized, "sale_amount")

    dates = [
        transaction_date
        for record in accepted
        for transaction_date in (record.purchase_date, record.sale_date)
        if transaction_date is not None
    ]
    reasons = Counter(
        validation.code
        for record in quarantined + ignored
        for validation in record.validation_results
    )
    rows_by_reason: dict[str, list[int]] = {}
    for record in quarantined + ignored:
        for validation in record.validation_results:
            rows_by_reason.setdefault(validation.code, []).append(record.source_row)
    return SheetReconciliation(
        source_sheet=sheet_name,
        company=company,
        source_rows_considered=len(data_records),
        valid_realized_lots=len(realized),
        provisionally_open_lots=len(open_lots),
        quarantined_rows=len(quarantined),
        ignored_rows=len(ignored),
        purchase_quantity=purchase_quantity,
        sale_quantity=sale_quantity,
        open_quantity=open_quantity,
        realized_quantity_difference=realized_purchase_quantity - sale_quantity,
        purchase_amount=purchase_amount,
        sale_amount=sale_amount,
        purchase_amount_difference=_amount_difference(accepted, "purchase"),
        sale_amount_difference=_amount_difference(realized, "sale"),
        earliest_transaction_date=min(dates) if dates else None,
        latest_transaction_date=max(dates) if dates else None,
        after_stated_period_rows=sum(
            any(
                validation.code == "after_stated_period" for validation in record.validation_results
            )
            for record in data_records
        ),
        reason_counts=tuple(sorted(reasons.items())),
        reason_rows=tuple(
            (code, tuple(row_numbers)) for code, row_numbers in sorted(rows_by_reason.items())
        ),
    )


def _sum_field(records: tuple[TradeLotRecord, ...], field: str) -> Decimal:
    return sum(
        (value for record in records if (value := getattr(record, field)) is not None),
        start=Decimal("0"),
    )


def _amount_difference(records: tuple[TradeLotRecord, ...], side: str) -> Decimal:
    differences: list[Decimal] = []
    for record in records:
        quantity = getattr(record, f"{side}_quantity")
        rate = getattr(record, f"{side}_rate")
        amount = getattr(record, f"{side}_amount")
        if quantity is not None and rate is not None and amount is not None:
            differences.append(amount - (quantity * rate))
    return sum(differences, start=Decimal("0"))


def render_ingestion_report(result: IngestionResult, *, workbook_name: str) -> str:
    """Render a concise Markdown reconciliation without exposing workbook credentials."""
    summaries = tuple(result.reconciliation_for(sheet) for sheet in _COMPANIES)
    lines = [
        "# Historical Trade Workbook Ingestion Report",
        "",
        f"Source workbook: `{workbook_name}` (opened read-only; decrypted in memory).",
        "",
        "Source rows considered are rows 4 through the last non-empty A:H row. "
        "The three validated header rows are retained as structural lineage but excluded here.",
        "",
        "| Sheet | Source rows | Realized | Open | Quarantined | Ignored |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for summary in summaries:
        lines.append(
            f"| {summary.source_sheet} | {summary.source_rows_considered} | "
            f"{summary.valid_realized_lots} | {summary.provisionally_open_lots} | "
            f"{summary.quarantined_rows} | {summary.ignored_rows} |"
        )

    for summary in summaries:
        reasons = ", ".join(f"{code}: {count}" for code, count in summary.reason_counts)
        reason_rows = "; ".join(
            f"`{code}`: rows {_format_row_numbers(row_numbers)}"
            for code, row_numbers in summary.reason_rows
        )
        lines.extend(
            [
                "",
                f"## {summary.source_sheet}",
                "",
                f"- Purchase quantity: {_format_quantity(summary.purchase_quantity)}",
                f"- Sale quantity: {_format_quantity(summary.sale_quantity)}",
                f"- Provisionally open quantity: {_format_quantity(summary.open_quantity)}",
                "- Realized quantity reconciliation difference: "
                f"{_format_quantity(summary.realized_quantity_difference)}",
                f"- Purchase amount: {_format_amount(summary.purchase_amount)}",
                f"- Sale amount: {_format_amount(summary.sale_amount)}",
                "- Purchase recorded-minus-calculated difference: "
                f"{_format_amount(summary.purchase_amount_difference)}",
                "- Sale recorded-minus-calculated difference: "
                f"{_format_amount(summary.sale_amount_difference)}",
                "- Transaction date range: "
                f"{_format_date(summary.earliest_transaction_date)} to "
                f"{_format_date(summary.latest_transaction_date)}",
                f"- Rows dated after 2025-03-31: {summary.after_stated_period_rows}",
                f"- Ignored/quarantined reasons: {reasons or 'none'}",
                f"- Ignored/quarantined source rows by reason: {reason_rows or 'none'}",
            ]
        )

    lines.extend(
        [
            "",
            "## Remaining limitations",
            "",
            "- This is an already FIFO-matched lot dataset, not an original execution ledger.",
            "- No portfolio NAV or execution-ledger completeness is claimed.",
            "- Cash balances, fees, taxes, dividends, corporate actions, and original order IDs "
            "are unavailable.",
            "- The workbook statement about unmatched sells remains unresolved and is not used "
            "as a holdings rule.",
            "- Repeated rows are retained; transactions after 2025-03-31 are retained and flagged.",
            f"- Recorded-versus-calculated amount tolerance: {AMOUNT_TOLERANCE} currency units.",
            "",
        ]
    )
    return "\n".join(lines)


def _format_quantity(value: Decimal) -> str:
    return f"{value:,.6f}".rstrip("0").rstrip(".")


def _format_amount(value: Decimal) -> str:
    rounded = value.quantize(Decimal("0.01"))
    return "0.00" if rounded == 0 else f"{rounded:,.2f}"


def _format_date(value: date | None) -> str:
    return value.isoformat() if value is not None else "n/a"


def _format_row_numbers(row_numbers: tuple[int, ...]) -> str:
    ranges: list[str] = []
    range_start = previous = row_numbers[0]
    for row_number in row_numbers[1:]:
        if row_number == previous + 1:
            previous = row_number
            continue
        ranges.append(str(range_start) if range_start == previous else f"{range_start}-{previous}")
        range_start = previous = row_number
    ranges.append(str(range_start) if range_start == previous else f"{range_start}-{previous}")
    return ", ".join(ranges)
